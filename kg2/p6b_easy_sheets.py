"""Easy-to-use annotation sheets, rebuilt from the sheets p6_validation_sheets.py made.

    python3 p6b_easy_sheets.py            # build into curation/_staging/ (nothing replaced)
    python3 p6b_easy_sheets.py --install  # archive the old sheets, put the new ones in curation/

Same items, same sheet names and same answer columns, so p6_score.py reads them unchanged; any
answer already typed is carried over. What changes for the person annotating:
  * a START HERE tab: what to do, one example per label, a live progress counter
  * the two entity names are bold in the sentence; PMID and entity IDs are clickable links
  * answer cells are yellow until filled; header frozen; rows sized to the sentence
  * the edge audit keeps only the 100 literature (L3) edges and shows the evidence sentence;
    the 200 curated/omics edges are checked by p6_v3_auto.py
  * the answer key moves to curation/_answer_key_do_not_open/
Never resamples: item IDs, sentences and order are copied from the existing sheets.
"""
import shutil
import sys
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import CURATION, OUT, VERSION, read_tsv

STAGE = CURATION / "_staging"
KEY = "V2_gold_set_KEY_do_not_open_before_annotation.xlsx"
KEY_DIR = CURATION / "_answer_key_do_not_open"
V2_LABELS = "positive_correlation,negative_correlation,association,binding,no_relation,cannot_tell"
YELLOW = PatternFill("solid", fgColor="FFFFF2B3")
HEAD = PatternFill("solid", fgColor="FFDDDDDD")
BOLD = InlineFont(b=True)


def rows_of(path, sheet):
    ws = load_workbook(path)[sheet]
    vals = list(ws.values)
    head = [str(h) for h in vals[0]]
    return [dict(zip(head, r)) for r in vals[1:] if any(v is not None for v in r)]


def id_url(curie):
    p, l = curie.split(":", 1)
    return {"NCBIGene": f"https://www.ncbi.nlm.nih.gov/gene/{l}",
            "MESH": f"https://meshb.nlm.nih.gov/record/ui?ui={l}"}.get(p)


def bold_entities(sentence, names):
    """Sentence as rich text with each entity name in bold (case-insensitive, first hit each)."""
    spans = []
    low = sentence.lower()
    for n in names:
        i = low.find((n or "").lower()) if n else -1
        if i >= 0:
            spans.append((i, i + len(n)))
    kept = []
    for a, b in sorted(spans):  # drop a span that overlaps the previous one
        if not kept or a >= kept[-1][1]:
            kept.append((a, b))
    spans = kept
    if not spans:
        return sentence
    out, pos = [], 0
    for a, b in spans:
        if a > pos:
            out.append(sentence[pos:a])
        out.append(TextBlock(BOLD, sentence[a:b]))
        pos = b
    if pos < len(sentence):
        out.append(sentence[pos:])
    return CellRichText(out)


def link(cell, text, url):
    cell.value = text
    if url:
        cell.hyperlink = url
        cell.font = Font(color="FF0563C1", underline="single")


def finish(ws, answer_cols, widths, n):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"  # printable: all columns on one page width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in answer_cols:  # yellow until answered
        ws.conditional_formatting.add(f"{col}2:{col}{n + 1}",
                                      FormulaRule(formula=[f'LEN({col}2)=0'], fill=YELLOW))


def start_sheet(wb, title, lines):
    ws = wb.create_sheet("START HERE", 0)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(lines, start=3):
        ws.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    return ws


def build_v2(src, dst, who):
    items = rows_of(src, "items")
    wb = Workbook()
    ws = wb.active
    ws.title = "items"
    cols = ["item", "sentence", "entity_a", "entity_b", "relation", "a_correct", "b_correct",
            "comment", "pmid", "id_a", "id_b"]
    ws.append(cols)
    for i, it in enumerate(items, start=2):
        ws.cell(i, 1, it["item"])
        ws.cell(i, 2).value = bold_entities(it["sentence"] or "", [it["entity_a"], it["entity_b"]])
        ws.cell(i, 3, it["entity_a"]).font = Font(bold=True)
        ws.cell(i, 4, it["entity_b"]).font = Font(bold=True)
        for j, c in enumerate(("relation", "a_correct", "b_correct", "comment"), start=5):
            ws.cell(i, j, it.get(c))
        link(ws.cell(i, 9), str(it["pmid"]), f"https://pubmed.ncbi.nlm.nih.gov/{it['pmid']}/")
        link(ws.cell(i, 10), it["id_a"], id_url(it["id_a"] or ":"))
        link(ws.cell(i, 11), it["id_b"], id_url(it["id_b"] or ":"))
        ws.row_dimensions[i].height = 15 * max(2, len(it["sentence"] or "") // 70 + 1)
    n = len(items)
    for col, opts in (("E", V2_LABELS), ("F", "yes,no,unsure"), ("G", "yes,no,unsure")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")
    finish(ws, ["E", "F", "G"], {"A": 6, "B": 70, "C": 18, "D": 18, "E": 22, "F": 10, "G": 10,
                                "H": 22, "I": 10, "J": 15, "K": 15}, n)
    st = start_sheet(wb, f"Gold set — {who} (work alone; do not look at the other annotator's file)", [
        "WHAT TO DO — for each row in the 'items' tab (400 rows, ~8 h; you can stop and continue any time):",
        "1. Read the sentence. Entity A and entity B are in the two bold columns next to it (the sentence may"
        " use another name for them, e.g. 'NRH:quinone oxidoreductase 2' for NQO2).",
        "2. 'relation' (yellow, pick from the list): does THE SENTENCE ITSELF state a relation between A and B?",
        "     positive_correlation — both go up together, or A increases/activates B.  e.g. 'GFAP rises with injury severity'",
        "     negative_correlation — A decreases/inhibits B, or opposite directions.  e.g. 'NQO2 inhibitor lowered ROS'",
        "     association — the sentence links them with no direction.  e.g. 'NfL was associated with outcome'",
        "     binding — they physically bind.  e.g. 'resveratrol binds NQO2'",
        "     no_relation — both named, but the sentence does not relate them",
        "     cannot_tell — too ambiguous to decide",
        "3. 'a_correct' / 'b_correct': is the ID right for that entity? Click the blue ID (columns J, K) to check."
        " yes / no / unsure.",
        "4. Optional comment. Yellow cells turn white when filled. Save with the same file name.",
        "",
        "Progress:",
    ])
    st["A16"] = '=COUNTA(items!E2:E401)&" of 400 relations done"'
    st["A16"].font = Font(bold=True, size=12)
    wb.save(dst)
    return sum(1 for it in items if it.get("relation"))


def build_v3(src, dst):
    edges = [e for e in rows_of(src, "edges") if e["layer"] == "L3"]
    ev = {}
    for r in read_tsv(OUT / VERSION / "edge_evidence.tsv"):
        if r["layer"] == "L3" and r["sentence"]:
            ev.setdefault(r["edge_id"], (r["pmid"], r["sentence"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "edges"
    cols = ["edge_id", "subject", "predicate", "object", "evidence_sentence", "supported",
            "comment", "pmid", "layer", "source"]
    ws.append(cols)
    for i, e in enumerate(edges, start=2):
        pmid, sent = ev.get(e["edge_id"], ((e["publications"] or "").split(";")[0], ""))
        ws.cell(i, 1, e["edge_id"])
        ws.cell(i, 2, e["subject"]).font = Font(bold=True)
        ws.cell(i, 3, (e["predicate"] or "").replace("biolink:", ""))
        ws.cell(i, 4, e["object"]).font = Font(bold=True)
        ws.cell(i, 5).value = bold_entities(sent, [e["subject"], e["object"]]) if sent else \
            "(no single sentence names both — open the PMID)"
        ws.cell(i, 6, e.get("supported"))
        ws.cell(i, 7, e.get("comment"))
        link(ws.cell(i, 8), str(pmid), f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/" if pmid else None)
        ws.cell(i, 9, e["layer"])
        ws.cell(i, 10, e["source"])
        ws.row_dimensions[i].height = 15 * max(2, len(sent) // 70 + 1)
    n = len(edges)
    dv = DataValidation(type="list", formula1='"yes,no,partly,cannot_check"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{n + 1}")
    finish(ws, ["F"], {"A": 15, "B": 18, "C": 28, "D": 22, "E": 70, "F": 13, "G": 22, "H": 10,
                       "I": 6, "J": 17}, n)
    st = start_sheet(wb, "Edge audit — literature edges (100 rows, ~1 h)", [
        "Each row is one edge text-mined from an abstract: subject — predicate — object.",
        "Read the evidence sentence next to the edge; click the PMID if you need the whole abstract.",
        "'supported' (yellow): does the source actually support THIS edge, with this predicate?",
        "     yes — clearly stated   ·   partly — related but the predicate/direction is off",
        "     no — not supported or wrong entities   ·   cannot_check — not enough text to judge",
        "The 200 curated and omics edges were checked automatically against their source files — not here.",
        "",
        "Progress:",
    ])
    st["A11"] = f'=COUNTA(edges!F2:F{n + 1})&" of {n} done"'
    st["A11"].font = Font(bold=True, size=12)
    wb.save(dst)
    return n


def main():
    install = "--install" in sys.argv
    STAGE.mkdir(exist_ok=True)
    kept = build_v2(CURATION / "V2_gold_set_annotator_1.xlsx", STAGE / "V2_gold_set_annotator_1.xlsx", "Annotator 1")
    kept2 = build_v2(CURATION / "V2_gold_set_annotator_2.xlsx", STAGE / "V2_gold_set_annotator_2.xlsx", "Annotator 2")
    n3 = build_v3(CURATION / "V3_edge_audit.xlsx", STAGE / "V3_edge_audit.xlsx")
    print(f"staged in {STAGE}: annotator_1 ({kept} answers carried over), annotator_2 ({kept2}), "
          f"edge audit {n3} L3 rows")
    if install:
        arch = CURATION / f"_archive_{date.today().isoformat()}"
        arch.mkdir(exist_ok=True)
        for f in ("V2_gold_set_annotator_1.xlsx", "V2_gold_set_annotator_2.xlsx", "V3_edge_audit.xlsx"):
            shutil.move(CURATION / f, arch / f)
            shutil.move(STAGE / f, CURATION / f)
        KEY_DIR.mkdir(exist_ok=True)
        if (CURATION / KEY).exists():
            shutil.move(CURATION / KEY, KEY_DIR / KEY)
        STAGE.rmdir()
        print(f"installed; originals in {arch}; answer key in {KEY_DIR}")


if __name__ == "__main__":
    main()
