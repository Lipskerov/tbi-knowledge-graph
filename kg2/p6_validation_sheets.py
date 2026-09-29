"""P6a — human validation sheets (V1-V3). Annotators fill them; p6_score.py scores them.

curation/V2_gold_set.xlsx   400 sentences: 300 PubTator3 relations (stratified by predicate)
                            + 100 co-mention pairs PubTator3 did NOT relate (for recall).
                            The extractor's label is HIDDEN; two annotators label blind.
curation/V3_edge_audit.xlsx 100 random edges per layer (L1, L2, L3): is the edge supported by
                            its cited source?
Fixed seed: re-running regenerates the identical sample (never resample after looking).
"""
import json
import random
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from common import CURATION, OUT, RAW, VERSION, read_tsv
from p4_literature import curie, sentences

SEED = 20260929
REL = OUT / VERSION
LABELS = "positive_correlation,negative_correlation,association,binding,no_relation,cannot_tell"


def sheet(wb, title, cols, rows, validations):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    for col_letter, options in validations:
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{len(rows) + 1}")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 60 if col[0].value == "sentence" else 22
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def negatives(n, rng):
    """Sentences where a Gene/Chemical/Disease pair co-occurs but PubTator3 gave no relation."""
    pool = []
    for path in sorted((RAW / "pubtator").glob("batch_*.json")):
        for doc in json.loads(path.read_text())["PubTator3"]:
            related = set()
            for rel in doc.get("relations", []):
                a, b = curie(rel["infons"].get("role1", {})), curie(rel["infons"].get("role2", {}))
                related |= {(a, b), (b, a)}
            for p in doc["passages"]:
                if p["infons"].get("type") not in ("title", "front", "abstract"):
                    continue
                anns = [(curie(a["infons"]), a["text"], a["locations"][0]["offset"])
                        for a in p["annotations"] if curie(a["infons"])]
                for s0, s1, s in sentences(p["text"], p["offset"]):
                    inside = {}
                    for c, txt, off in anns:
                        if s0 <= off < s1:
                            inside.setdefault(c, txt)
                    ids = sorted(inside)
                    for i in range(len(ids)):
                        for j in range(i + 1, len(ids)):
                            if (ids[i], ids[j]) not in related:
                                pool.append({"pmid": doc["pmid"], "sentence": s,
                                             "entity_a": inside[ids[i]], "id_a": ids[i],
                                             "entity_b": inside[ids[j]], "id_b": ids[j],
                                             "_truth": "none"})
    return rng.sample(pool, n)


def main():
    rng = random.Random(SEED)
    nodes = {n["id"]: n for n in read_tsv(REL / "nodes.tsv")}
    edges = {e["id"]: e for e in read_tsv(REL / "edges.tsv")}
    ev = [r for r in read_tsv(REL / "edge_evidence.tsv") if r["layer"] == "L3" and r["sentence"]]
    by_pred = defaultdict(list)
    for r in ev:
        by_pred[edges[r["edge_id"]]["predicate"]].append(r)
    quota = {"biolink:associated_with": 120, "biolink:positively_correlated_with": 80,
             "biolink:negatively_correlated_with": 80, "biolink:physically_interacts_with": 12,
             "biolink:related_to": 8}
    pos = []
    for pred, k in quota.items():
        for r in rng.sample(by_pred[pred], min(k, len(by_pred[pred]))):
            e = edges[r["edge_id"]]
            pos.append({"pmid": r["pmid"], "sentence": r["sentence"],
                        "entity_a": nodes[e["subject"]]["name"], "id_a": e["subject"],
                        "entity_b": nodes[e["object"]]["name"], "id_b": e["object"],
                        "_truth": pred})
    items = pos + negatives(100, rng)
    rng.shuffle(items)
    for i, it in enumerate(items):
        it["item"] = f"G{i:03d}"
    key = [{"item": it["item"], "pubtator3_label": it["_truth"], "id_a": it["id_a"],
            "id_b": it["id_b"]} for it in items]

    for annot in ("annotator_1", "annotator_2"):
        wb = Workbook()
        wb.remove(wb.active)
        guide = wb.create_sheet("README")
        for line in [
            "V2 gold set — label BLIND. Do not look at the other annotator's sheet.",
            "For each sentence: does the SENTENCE ITSELF state a relation between entity A and entity B?",
            "relation: positive_correlation (both go up together / A increases B), "
            "negative_correlation (A decreases B / inverse), association (linked, no sign), "
            "binding (physical binding), no_relation, cannot_tell.",
            "a_correct / b_correct: is the entity ID right for the text? Click the ID "
            "(NCBIGene → ncbi.nlm.nih.gov/gene, MESH → meshb.nlm.nih.gov).",
            "Budget: ~400 items, about 8 hours. Save as the same file name.",
        ]:
            guide.append([line])
        sheet(wb, "items", ["item", "pmid", "sentence", "entity_a", "id_a", "entity_b", "id_b",
                            "relation", "a_correct", "b_correct", "comment"], items,
              [("H", LABELS), ("I", "yes,no,unsure"), ("J", "yes,no,unsure")])
        wb.save(CURATION / f"V2_gold_set_{annot}.xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "key", ["item", "pubtator3_label", "id_a", "id_b"], key, [])
    wb.save(CURATION / "V2_gold_set_KEY_do_not_open_before_annotation.xlsx")

    audit = []
    for layer in ("L1", "L2", "L3"):
        pool = [e for e in edges.values() if e["layer"] == layer]
        for e in rng.sample(pool, 100):
            audit.append({"edge_id": e["id"], "layer": layer, "subject": nodes[e["subject"]]["name"],
                          "predicate": e["predicate"], "object": nodes[e["object"]]["name"],
                          "direction": e["direction"], "source": e["primary_knowledge_source"],
                          "publications": e["publications"][:120], "subject_id": e["subject"],
                          "object_id": e["object"]})
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "edges", ["edge_id", "layer", "subject", "predicate", "object", "direction", "source",
                        "publications", "subject_id", "object_id", "supported", "comment"], audit,
          [("K", "yes,no,partly,cannot_check")])
    wb.save(CURATION / "V3_edge_audit.xlsx")
    print(f"P6a: V2 items {len(items)} ({len(pos)} extractor relations + 100 negatives) x 2 annotators; "
          f"V3 audit {len(audit)} edges")


if __name__ == "__main__":
    main()
